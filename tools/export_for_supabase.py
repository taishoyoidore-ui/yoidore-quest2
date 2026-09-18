import os
import json
import openpyxl

def generate_supabase_stores_sql():
    excel_path = os.path.join(os.path.dirname(__file__), 'STORES.xlsx')
    output_sql_path = os.path.join(os.path.dirname(__file__), 'doc', 'stores_insert.sql')

    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found.")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if '店舗一覧' not in wb.sheetnames:
        print("Error: '店舗一覧' sheet not found.")
        return

    ws = wb['店舗一覧']
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Error: Sheet is empty.")
        return

    headers = [str(c or '').strip() for c in rows[0]]
    
    # Identify key columns flexibly
    def find_idx(*candidates):
        for c in candidates:
            c_low = c.lower()
            for idx, h in enumerate(headers):
                if h.lower() == c_low:
                    return idx
        return -1

    id_idx = find_idx('id', '店舗id')
    name_idx = find_idx('店舗名', '店名', '名')
    area_idx = find_idx('エリア', '地域')
    coupon_target_idx = find_idx('クーポン対象', 'クーポン対象店舗', '特典対象')

    sql_statements = [
        "-- STORES.xlsx から自動生成された店舗マスタ登録SQL",
        "TRUNCATE TABLE stores CASCADE;",
        "INSERT INTO stores (id, name, area, is_coupon_target, display_order, raw_data) VALUES"
    ]

    values_list = []
    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == '' for cell in row):
            continue

        store_name = str(row[name_idx] or '').strip() if (name_idx != -1 and name_idx < len(row)) else ''
        if not store_name:
            continue

        store_id = str(row[id_idx] or '').strip() if (id_idx != -1 and id_idx < len(row) and row[id_idx]) else f"store-{str(row_idx-1).zfill(2)}"
        area = str(row[area_idx] or '').strip() if (area_idx != -1 and area_idx < len(row) and row[area_idx]) else ''
        
        # Coupon target check
        is_coupon = False
        if coupon_target_idx != -1 and coupon_target_idx < len(row) and row[coupon_target_idx]:
            val = str(row[coupon_target_idx]).strip().lower()
            is_coupon = val in ('true', '1', 'yes', '〇', 'o', '可')
        else:
            is_coupon = (row_idx % 2 == 0)

        # Pack all columns into raw_data dictionary dynamically
        raw_data = {}
        for h_idx, h_name in enumerate(headers):
            if not h_name:
                continue
            val = row[h_idx] if h_idx < len(row) else None
            if val is not None:
                if hasattr(val, 'strftime'):
                    val = val.strftime('%Y-%m-%d %H:%M:%S')
                elif isinstance(val, str):
                    # Normalize carriage returns and newlines
                    val = val.replace('\r\n', '\n').replace('\r', '\n').strip()
                raw_data[h_name] = val

        raw_data_json = json.dumps(raw_data, ensure_ascii=False)
        escaped_name = store_name.replace("'", "''")
        escaped_area = area.replace("'", "''")
        is_coupon_str = 'true' if is_coupon else 'false'

        # Use dollar-quoting $json$...$json$ to guarantee zero syntax errors with quotes/newlines
        values_list.append(
            f"('{store_id}', '{escaped_name}', '{escaped_area}', {is_coupon_str}, {row_idx-1}, $json${raw_data_json}$json$::jsonb)"
        )

    sql_statements.append(",\n".join(values_list) + ";")

    with open(output_sql_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(sql_statements))

    print(f"Successfully generated {output_sql_path} with {len(values_list)} stores using dollar-quoted JSON.")

if __name__ == '__main__':
    generate_supabase_stores_sql()
